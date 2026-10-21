"""Строит эталонный solution/data/orders.xlsx: Orders + sentiment + Summary.

Запуск из каталога задачи:  python tools/solve_orders.py
"""
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

TASK = Path(__file__).resolve().parents[1]
LEXICON = {
    "all is good": "positive", "great service": "positive",
    "fast delivery, thanks": "positive", "product works perfectly": "positive",
    "very satisfied": "positive", "excellent quality": "positive",
    "will order again": "positive",
    "all is bad": "negative", "terrible support": "negative",
    "item arrived broken": "negative", "too expensive for this quality": "negative",
    "very disappointed": "negative", "slow shipping": "negative",
    "would not recommend": "negative",
    "where is my invoice": "neutral", "please confirm delivery date": "neutral",
    "do you have this in black": "neutral", "how to return the item": "neutral",
    "tracking number not received": "neutral", "is warranty included": "neutral",
}


def canon(text):
    return " ".join(str(text).split()).casefold()


def main():
    env = TASK / "env" / "data" / "orders.xlsx"
    out = TASK / "solution" / "data" / "orders.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(env, out)

    wb = load_workbook(out)
    ws = wb["Orders"]
    ws.cell(row=1, column=7).value = "sentiment"

    counts = Counter()
    r = 2
    while ws.cell(row=r, column=1).value is not None:
        fb = canon(ws.cell(row=r, column=6).value)
        sentiment = LEXICON[fb]
        ws.cell(row=r, column=7).value = sentiment
        counts[sentiment] += 1
        r += 1

    summary = wb.create_sheet("Summary")
    summary.append(["sentiment", "count"])
    for cls in ("positive", "negative", "neutral"):
        summary.append([cls, counts[cls]])

    wb.save(out)
    print(f"wrote {out} (rows={r - 2}, counts={dict(counts)})")


if __name__ == "__main__":
    main()
