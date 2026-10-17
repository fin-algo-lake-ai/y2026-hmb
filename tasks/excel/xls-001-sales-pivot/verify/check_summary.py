"""Verifier для xls-001-sales-pivot (нужен openpyxl).

Считает эталонную сводку из исходного env/data/orders.xlsx и сверяет лист
"Summary" в workspace. Печатает по строке на проверку и в конце
"SCORE: <доля пройденных>". Код выхода 0 — verifier отработал корректно.
"""
import os
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook

WS = Path(os.environ["HMB_WORKSPACE"])
WB = WS / "data" / "orders.xlsx"
PRISTINE = Path(__file__).resolve().parents[1] / "env" / "data" / "orders.xlsx"
TOL = 0.011


def close(a, b):
    try:
        return abs(float(a) - float(b)) <= TOL
    except (TypeError, ValueError):
        return False


def expected_summary():
    ws = load_workbook(PRISTINE, data_only=True)["Orders"]
    agg = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        _, region, _product, units, price = row
        u, rev = agg.get(region, (0, 0.0))
        agg[region] = (u + units, rev + units * price)
    out = OrderedDict()
    for region in sorted(agg):
        out[region] = agg[region]
    return out


def main():
    checks = []

    def add(name, ok, hint=""):
        checks.append((name, bool(ok), hint))

    wsv = load_workbook(WB, data_only=True)
    wsf = load_workbook(WB, data_only=False)

    exp_orders = [tuple(r) for r in
                  load_workbook(PRISTINE, data_only=True)["Orders"].iter_rows(values_only=True)]
    got_orders = ([tuple(r) for r in wsv["Orders"].iter_rows(values_only=True)]
                  if "Orders" in wsv.sheetnames else None)
    add("orders_sheet_intact", got_orders == exp_orders,
        "лист Orders отсутствует или изменён")

    exp = expected_summary()
    regions = list(exp)
    has = "Summary" in wsv.sheetnames
    add("summary_sheet_exists", has, "нет листа Summary")

    if has:
        s, sf = wsv["Summary"], wsf["Summary"]
        hdr = tuple(s.cell(row=1, column=c).value for c in (1, 2, 3))
        add("headers", hdr == ("region", "units", "revenue"), f"A1:C1 = {hdr}")

        got_regions = [s.cell(row=2 + i, column=1).value for i in range(len(regions))]
        no_extra = s.cell(row=2 + len(regions), column=1).value is None
        add("region_rows", got_regions == regions and no_extra,
            f"got={got_regions}, после последней строки пусто: {no_extra}")

        for i, region in enumerate(regions):
            r = 2 + i
            units_exp, rev_exp = exp[region]
            v_u = s.cell(row=r, column=2).value
            v_r = s.cell(row=r, column=3).value
            f_r = sf.cell(row=r, column=3).value
            hint = ""
            if v_r is None and isinstance(f_r, str) and f_r.startswith("="):
                hint = "формула без кэша значений — запиши вычисленное число"
            add(f"values[{region}]",
                close(v_u, units_exp) and close(v_r, rev_exp),
                hint or f"got=({v_u!r}, {v_r!r}), "
                        f"expected=({units_exp}, {round(rev_exp, 2)})")
    else:
        add("headers", False, "нет листа Summary")
        add("region_rows", False, "нет листа Summary")
        for region in regions:
            add(f"values[{region}]", False, "нет листа Summary")

    passed = sum(1 for _, ok, _ in checks if ok)
    for name, ok, hint in checks:
        print(("ok   " if ok else "FAIL ") + name + (f" ({hint})" if hint and not ok else ""))
    print(f"SCORE: {passed / len(checks):.2f}")


if __name__ == "__main__":
    main()
