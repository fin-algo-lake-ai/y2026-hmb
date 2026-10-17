"""Строит эталонный solution/data/orders.xlsx: лист Orders + лист Summary.

Запуск из каталога задачи:  python tools/solve_orders.py
"""
from pathlib import Path

from openpyxl import load_workbook

TASK = Path(__file__).resolve().parents[1]

wb = load_workbook(TASK / "env" / "data" / "orders.xlsx")
ws = wb["Orders"]

agg = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    _, region, _product, units, price = row
    u, rev = agg.get(region, (0, 0.0))
    agg[region] = (u + units, rev + units * price)

summary = wb.create_sheet("Summary")
summary.append(["region", "units", "revenue"])
for region in sorted(agg):
    u, rev = agg[region]
    summary.append([region, u, round(rev, 2)])

out = TASK / "solution" / "data" / "orders.xlsx"
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"wrote {out}")
